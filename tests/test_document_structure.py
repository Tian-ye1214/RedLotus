import io
import json

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from redlotus.tools.references import DocumentReader


async def test_ppt_group_images_keep_distinct_contents(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for colour in ("red", "blue"):
        stream = io.BytesIO()
        Image.new("RGB", (8, 8), colour).save(stream, format="PNG")
        stream.seek(0)
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(stream, Inches(1), Inches(1))
    path = tmp_path / "groups.pptx"
    presentation.save(path)
    parts = await DocumentReader().read(path, tmp_path)
    images = [p for p in parts if p.kind == "image"]
    assert [Image.open(p.path).getpixel((0, 0)) for p in images] == [
        (255, 0, 0),
        (0, 0, 255),
    ]


async def test_csv_quoted_semicolon_columns_are_preserved(tmp_path):
    path = tmp_path / "regions.csv"
    path.write_text('region;total\n"A;zone";2\nB;3\n', encoding="utf-8")
    parts = await DocumentReader().read(path, tmp_path)
    assert json.loads(parts[0].text) == [
        ["region", "total"],
        ["A;zone", "2"],
        ["B", "3"],
    ]


async def test_html_retains_table_relationships_and_link_targets(tmp_path):
    path = tmp_path / "report.html"
    path.write_text(
        '<h1>报告</h1><table><tr><th>区域</th><th>金额</th></tr><tr><td>北</td><td>12</td></tr></table><p><a href="https://example.test/source">来源</a></p><script>do_not_include()</script>',
        encoding="utf-8",
    )
    parts = await DocumentReader().read(path, tmp_path)
    tables = [p for p in parts if "表格" in p.locator]
    assert len(tables) == 1
    assert json.loads(tables[0].text) == [["区域", "金额"], ["北", "12"]]
    text = "\n".join(p.text for p in parts)
    assert "https://example.test/source" in text and "do_not_include" not in text
