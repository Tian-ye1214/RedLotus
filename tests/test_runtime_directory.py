import asyncio
import errno
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from pydantic_ai import BinaryContent

from redlotus.runtime import config, resources
from redlotus.runtime.network import ModelInputPolicy
from redlotus.prompts.prompt import get_skills_layout_text
from redlotus.tools.base_tools import BasicToolkit
from redlotus.tools.execution import get_execution_environment
from redlotus.tools.references import DocumentReader, OfficeConverter, ReferenceStore
from redlotus.tools.registry import SkillsManager, resolve_readable_path


@pytest.fixture
def workspace(tmp_path, isolated_config):
    return resources.WorkspaceContext.from_path(tmp_path)


@pytest.mark.parametrize('configured', [None, '', '   '])
@pytest.mark.asyncio
async def test_chat_files_and_bundled_skills_without_runtime(workspace, isolated_config, configured):
    isolated_config['storage']['runtime_dir'] = configured
    manager = SkillsManager(workspace=workspace)
    toolkit = BasicToolkit(manager, workspace=workspace, show_diff=lambda *a, **kw: (0, 0, 0))
    assert 'Saved' in await toolkit.write_file('WorkDatabase/notes.txt', content='hello')
    assert toolkit.read_file('WorkDatabase/notes.txt').return_value == 'hello'
    assert manager.skills
    layout = get_skills_layout_text(manager)
    assert str(resources.skills_dir()) in layout
    assert 'storage.runtime_dir' in layout
    assert resolve_readable_path('skills', work_base=workspace.root) == resources.skills_dir()


def test_missing_bundled_skill_read_without_runtime_reports_missing_file(workspace, isolated_config):
    isolated_config['storage']['runtime_dir'] = None
    toolkit = BasicToolkit(SkillsManager(workspace=workspace), workspace=workspace, show_diff=lambda *a, **kw: (0, 0, 0))
    result = toolkit.read_file('skills/absent.txt')
    assert result.startswith("Error reading 'skills/absent.txt':")
    assert 'storage.runtime_dir' not in result


def test_excel_reader_closes_first_workbook_when_second_open_fails(tmp_path, monkeypatch):
    import openpyxl

    source = tmp_path / 'source'
    source.write_bytes(b'not used by the injected loader')
    first = SimpleNamespace(closed=False)
    first.close = lambda: setattr(first, 'closed', True)
    calls = 0

    def load_workbook(stream, *, read_only, data_only):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ValueError('second workbook rejected')
        return first

    monkeypatch.setattr(openpyxl, 'load_workbook', load_workbook)
    with pytest.raises(ValueError, match='second workbook rejected'):
        DocumentReader().excel(source, tmp_path)
    assert first.closed


@pytest.mark.parametrize('operation', ['execution', 'install', 'office'])
def test_required_runtime_diagnostic(workspace, operation):
    with resources.workspace_context(workspace), pytest.raises(config.ConfigError) as error:
        if operation == 'execution':
            get_execution_environment(workspace=workspace)
        elif operation == 'install':
            resources.user_skills_dir(workspace)
        else:
            asyncio.run(OfficeConverter().convert(workspace.root / 'old.doc', 'docx', workspace.root / 'out'))
    message = str(error.value)
    assert 'storage.runtime_dir' in message
    assert 'WorkDatabase/runtime' in message
    assert '作用' in message


@pytest.mark.parametrize('name', ['runtime_dir', 'project_dir', 'references_dir'])
def test_storage_rejects_escape(workspace, isolated_config, name):
    isolated_config['storage'][name] = '../outside'
    with pytest.raises(config.ConfigError, match='当前项目'):
        resources._project_storage_path(name, workspace)


@pytest.mark.parametrize('second_name,second_mime', [('second.txt', 'text/plain'), ('first.txt', 'audio/wav')])
def test_identical_bytes_in_same_input_slot_keep_name_and_mime(workspace, second_name, second_mime):
    async def run():
        store = ReferenceStore(workspace)
        policy = ModelInputPolicy(8, 1024)
        first = await store.import_binary(BinaryContent(data=b'hello', media_type='text/plain', identifier='first.txt'), source='attachment:0', policy=policy)
        second = await store.import_binary(BinaryContent(data=b'hello', media_type=second_mime, identifier=second_name), source='attachment:0', policy=policy)
        assert first.id != second.id
        assert (second.name, second.media_type) == (second_name, second_mime)
        assert (store.load(second.id).name, store.load(second.id).media_type) == (second_name, second_mime)
        assert second.parts[0].kind == ('audio' if second_mime.startswith('audio/') else 'text')
        assert first.snapshot.read_bytes() == second.snapshot.read_bytes() == b'hello'
    asyncio.run(run())


def test_optional_runtime_keeps_layered_precedence(workspace, monkeypatch):
    sources = ({'storage': {'runtime_dir': ' '}}, {'storage': {'runtime_dir': 'cache/runtime'}}, {'storage': {'runtime_dir': 'fallback'}})
    monkeypatch.setattr(resources, 'settings', lambda: config.ConfigValues(sources=sources))
    assert resources.runtime_dir(workspace) == workspace.root / 'cache/runtime'


def test_readable_paths_are_bound_to_supplied_workspace(workspace, isolated_config, tmp_path):
    isolated_config['storage']['runtime_dir'] = 'runtime'
    with resources.workspace_context(resources.WorkspaceContext.from_path(tmp_path / 'other')):
        assert resolve_readable_path('skills/new/SKILL.md', work_base=workspace.root) == workspace.root / 'runtime/skills/new/SKILL.md'


@pytest.mark.parametrize('command', ['echo hello', 'npx clawhub install example', 'npx clawhub --dir skills install example'])
def test_command_missing_runtime_returns_actionable_error(workspace, command):
    toolkit = BasicToolkit(SkillsManager(workspace=workspace), workspace=workspace, show_diff=lambda *a, **kw: (0, 0, 0))
    result = asyncio.run(toolkit.run_command(command))
    assert 'storage.runtime_dir' in result
    assert 'WorkDatabase/runtime' in result


def test_cleanup_does_not_remove_persistent_skills(workspace, isolated_config):
    isolated_config['storage']['runtime_dir'] = 'runtime'
    cache = resources.runtime_dir(workspace) / 'cache'
    owned = cache / 'inactive'
    owned.mkdir(parents=True)
    (owned / 'cache.bin').write_bytes(b'cache')
    skill = resources.user_skills_dir(workspace) / 'custom/SKILL.md'
    skill.parent.mkdir(parents=True)
    skill.write_text('persistent', encoding='utf-8')
    assert resources._delete_owned_cache(cache, owned) == 5
    assert skill.read_text(encoding='utf-8') == 'persistent'


def test_recorded_tree_removal_requires_exact_plain_contents(tmp_path):
    root = tmp_path / 'owned'
    (root / 'nested').mkdir(parents=True)
    (root / 'nested/file.bin').write_bytes(b'model')
    marker = {'schema': 1, 'owner': 'test'}
    resources.atomic_write_json(root / '.marker.json', marker)
    extra = root / 'user-note.txt'
    extra.write_text('keep')
    assert not resources.remove_recorded_tree(root, '.marker.json', marker, ['nested/file.bin'], complete=True)
    assert extra.exists()
    extra.unlink()
    assert resources.remove_recorded_tree(root, '.marker.json', marker, ['nested/file.bin'], complete=True)
    assert not root.exists()
    with pytest.raises(ValueError):
        resources.owned_path(tmp_path, '../escape')


@pytest.mark.parametrize('name,allowed', [('.', False), ('root/', True), ('root/file', True),
                                           ('root/../escape', False), ('root\\file', False)])
def test_safe_tar_name(name, allowed):
    assert resources.safe_tar_name(name, 'root') is allowed


@pytest.mark.skipif(resources.os.name != 'nt', reason='Windows sharing retry')
def test_atomic_write_retries_temporary_sharing_error(tmp_path, monkeypatch):
    destination = tmp_path / 'atomic.json'
    real_replace = resources.os.replace
    denied = []
    def replace(source, target):
        if target == destination and not denied:
            denied.append(True)
            error = PermissionError(13, 'temporary sharing error')
            error.winerror = 5
            raise error
        return real_replace(source, target)
    monkeypatch.setattr(resources.os, 'replace', replace)
    resources.atomic_write_json(destination, {'ok': True})
    assert denied and json.loads(destination.read_text()) == {'ok': True}


@pytest.mark.skipif(resources.os.name != 'nt', reason='Windows sharing retry')
def test_replace_retry_moves_directory_after_temporary_denial(tmp_path, monkeypatch):
    source, destination = tmp_path / 'source', tmp_path / 'destination'
    source.mkdir()
    (source / 'model.bin').write_bytes(b'weight')
    real_replace = resources.os.replace
    denied = []
    def replace(first, second):
        if second == destination and not denied:
            denied.append(True)
            error = PermissionError(13, 'temporary directory denial')
            error.winerror = 5
            raise error
        return real_replace(first, second)
    monkeypatch.setattr(resources.os, 'replace', replace)
    resources.replace_retry(source, destination)
    assert denied and not source.exists() and (destination / 'model.bin').read_bytes() == b'weight'


def test_cleanup_missing_optional_runtime_preserves_disk_full_error(workspace, isolated_config):
    from redlotus.sessions.cleanup import retry_after_storage_cleanup
    isolated_config['storage']['cleanup'] = {'enabled': True, 'execution_cache': True, 'session_retention_days': 1}
    error = OSError(errno.ENOSPC, 'disk full')
    with resources.workspace_context(workspace), pytest.raises(OSError) as caught:
        retry_after_storage_cleanup(workspace.root / 'out', error, lambda: None)
    assert caught.value is error


def test_reference_concurrent_imports_are_immutable(workspace):
    async def run():
        store = ReferenceStore(workspace)
        policy = ModelInputPolicy(8, 1024)
        refs = await asyncio.gather(*(
            store.import_bytes(b'hello', name=name, source='attachment:0', media_type='text/plain', policy=policy)
            for name in ['first.txt', 'second.txt', 'first.txt']
        ))
        assert refs[0] == refs[2]
        assert refs[0].id != refs[1].id
        assert refs[0].snapshot == refs[1].snapshot
        assert all(store.load(ref.id) == ref for ref in refs)
        assert all(ref.parts[0].text == 'hello' for ref in refs)
    asyncio.run(run())


@pytest.mark.parametrize('protected', ['active_turn', 'interrupted_turn', 'paused_turn', 'tasks', 'lock', None])
def test_session_cleanup_preserves_active_and_locked_journals(workspace, protected):
    import time
    from contextlib import nullcontext
    from filelock import FileLock
    from redlotus.sessions.cleanup import _delete_old_session
    from redlotus.sessions.storage import SessionFile
    session = SessionFile.create(resources.session_data_dir(workspace), workspace.project_id, session_id='session', workspace=workspace)
    if protected and protected != 'lock':
        session.update(metadata={protected: [{'status': 'running'}] if protected == 'tasks' else {'turn_id': 'pending'}})
    path = session.path
    lock = FileLock(path.parent / '.use-test.lock') if protected == 'lock' else nullcontext()
    with lock:
        result = _delete_old_session(path.parent, path, time.time() + 10, workspace.project_id)
    assert path.exists() is bool(protected)
    assert (result is None) is bool(protected)


def test_cleanup_reclaims_only_owned_cache_and_retries(workspace, isolated_config):
    from redlotus.sessions.cleanup import retry_after_storage_cleanup
    isolated_config['storage'].update(runtime_dir='runtime', cleanup={'enabled': True, 'execution_cache': True, 'session_retention_days': 1})
    resources.session_data_dir(workspace).mkdir(parents=True)
    runtime = resources.runtime_dir(workspace)
    inactive = runtime / 'cache/inactive'
    inactive.mkdir(parents=True)
    (inactive / '.redlotus-cache').write_text(json.dumps({'project_id': 'inactive'}), encoding='utf-8')
    skill = runtime / 'skills/installed/SKILL.md'
    skill.parent.mkdir(parents=True)
    skill.write_text('persistent', encoding='utf-8')
    retried = []
    with resources.workspace_context(workspace):
        retry_after_storage_cleanup(workspace.root / 'out', OSError(errno.ENOSPC, 'disk full'), lambda: retried.append(True))
    assert retried == [True]
    assert not inactive.exists()
    assert skill.read_text(encoding='utf-8') == 'persistent'


@pytest.mark.asyncio
async def test_file_operations_preserve_text_and_review_decisions(workspace):
    from redlotus.tools.base_tools import reconstruct
    toolkit = BasicToolkit(SkillsManager(workspace=workspace), workspace=workspace, show_diff=lambda *a, **kw: (0, 0, 0))
    toolkit.review_store.activate(lambda: None)
    await toolkit.write_file('WorkDatabase/one.txt', content='first\r\nsecond\r\nthird\r\n')
    await toolkit.write_file('WorkDatabase/copy.txt', copy_from='WorkDatabase/one.txt')
    assert toolkit.read_file('WorkDatabase/copy.txt').return_value == 'first\r\nsecond\r\nthird\r\n'
    assert 'copy.txt' in toolkit.list_files('WorkDatabase')
    assert 'copy.txt:2: second' in toolkit.search_in_files('SECOND', '.txt')
    await toolkit.edit_file('WorkDatabase/one.txt', 'first', 'FIRST')
    entry = next(e for e in toolkit.review_store.entries() if e.name == 'WorkDatabase/one.txt')
    assert toolkit.review_store.decide(entry, 0, True)
    assert not (workspace.root / 'WorkDatabase/one.txt').exists()
    assert reconstruct('a\nb\nc\nd\ne\n', 'A\nb\nc\nd\nE\n', {0}) == 'a\nb\nc\nd\nE\n'


@pytest.mark.parametrize('async_tool', [False, True])
@pytest.mark.parametrize('fails', [False, True])
def test_tool_telemetry_records_result_once(monkeypatch, async_tool, fails):
    from redlotus.tools import registry
    events = []
    monkeypatch.setattr(registry, '_notify', lambda *args: None)
    monkeypatch.setattr(registry.TRACE_STORE, 'record', lambda *args, **kwargs: events.append(kwargs))
    def operation():
        if fails:
            raise ValueError('broken')
        return 'done'
    async def async_operation():
        return operation()
    wrapped = registry._wrap(async_operation if async_tool else operation)
    try:
        result = asyncio.run(wrapped()) if async_tool else wrapped()
        assert result == 'done'
    except ValueError:
        assert fails
    assert len(events) == 1
    assert events[0]['success'] is not fails
