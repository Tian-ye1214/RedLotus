from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import shutil
import tempfile
from pathlib import Path

from redlotus.infra.paths import user_data_dir
from redlotus.config.app_config import get_env
from redlotus.infra.subprocess_runner import run_subprocess
from redlotus.infra.persist_utils import finish_file_io
from redlotus.references.models import ReferencePart


class OfficeConverter:
    """A private headless Office profile, with cancellable process-tree ownership."""

    @staticmethod
    def executable() -> str:
        candidates = [
            get_env("LIBREOFFICE_PATH", warn=False),
            shutil.which("soffice.com"),
            shutil.which("soffice"),
            str(
                Path(os.environ.get("ProgramFiles", "C:/Program Files"))
                / "LibreOffice/program/soffice.com"
            ),
            str(user_data_dir() / "tools/libreoffice/program/soffice.com"),
        ]
        for value in candidates:
            if value and Path(value).is_file():
                return value
        raise ValueError(
            "DOC/PPT 转换需要 LibreOffice；请安装或设置 LIBREOFFICE_PATH。"
        )

    async def convert(self, source: Path, target_format: str, directory: Path) -> Path:
        executable = self.executable()
        directory.mkdir(parents=True, exist_ok=True)
        # LibreOffice still uses Windows APIs with limited path lengths; keep its private
        # profile and working copies short, then publish the result to the reference store.
        with tempfile.TemporaryDirectory(prefix="rl-office-") as profile:
            profile_root = Path(profile)
            user = profile_root / "user"
            user.mkdir()
            working = profile_root / "document" / source.name
            working.parent.mkdir()
            shutil.copyfile(source, working)
            (user / "registrymodifications.xcu").write_text(
                '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
                '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>',
                encoding="utf-8",
            )
            args = [
                executable,
                "-env:UserInstallation=" + profile_root.as_uri(),
                "--headless",
                "--nologo",
                "--nodefault",
                "--norestore",
                "--convert-to",
                target_format,
                "--outdir",
                str(profile_root),
                str(working),
            ]
            stdout, stderr, code = await run_subprocess(
                args, shell=False, cwd=str(profile_root), timeout=120
            )
            output = profile_root / (source.stem + "." + target_format.split(":")[0])
            if code != 0 or not output.is_file():
                raise ValueError(
                    f"Office 转换失败 ({source.name}, exit {code})：{stdout}\n{stderr}"
                )
            target = directory / output.name
            shutil.copyfile(output, target)
            return target


class DocumentReader:
    """Read format structure once; never execute referenced content."""

    async def read(self, source: Path, directory: Path) -> list[ReferencePart]:
        extension = source.suffix.lower()
        if extension in (".doc", ".ppt", ".xls"):
            modern = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}[extension]
            source = await OfficeConverter().convert(
                source, modern, directory / "converted"
            )
            extension = "." + modern
        readers = {
            ".pdf": self.pdf,
            ".docx": self.word,
            ".pptx": self.slides,
            ".xlsx": self.excel,
            ".csv": self.csv,
            ".html": self.html,
            ".htm": self.html,
        }
        return await finish_file_io(
            asyncio.to_thread(readers.get(extension, self.text), source, directory)
        )

    @staticmethod
    def decode(data: bytes) -> str:
        for encoding in (
            "utf-8-sig",
            "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "gb18030",
        ):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                pass
        raise ValueError("无法识别引用文件的文本编码。")

    def text(self, source: Path, directory: Path) -> list[ReferencePart]:
        text = self.decode(source.read_bytes())
        if "\x00" in text:
            raise ValueError(f"没有配置此二进制格式的解析器：{source.suffix}")
        return [ReferencePart.from_text(text)]

    def csv(self, source: Path, directory: Path) -> list[ReferencePart]:
        text = self.decode(source.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        return [ReferencePart.from_text(rows, locator="CSV 行列")]

    def html(self, source: Path, directory: Path) -> list[ReferencePart]:
        from lxml import html, etree

        root = html.fromstring(
            self.decode(source.read_bytes()).encode("utf-8"),
            parser=html.HTMLParser(encoding="utf-8"),
        )
        etree.strip_elements(root, "script", "style", with_tail=False)
        tables = []
        for number, table in enumerate(root.xpath("self::table | .//table"), 1):
            rows = [
                [cell.text_content().strip() for cell in row.xpath("./th | ./td")]
                for row in table.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
            ]
            tables.append(ReferencePart.from_text(rows, locator=f"HTML 表格 {number}"))
            table.drop_tree()
        for link in root.xpath(".//a[@href]"):
            link.tail = f" ({link.get('href')})" + (link.tail or "")
        for element in root.iter():
            if element.tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
                element.text = "#" * int(element.tag[1]) + " " + (element.text or "")
            if element.tag in (
                "p",
                "div",
                "section",
                "li",
                "br",
                "h1",
                "h2",
                "h3",
                "h4",
                "h5",
                "h6",
            ):
                element.tail = "\n" + (element.tail or "")
        return [
            ReferencePart.from_text(root.text_content().strip(), locator="HTML 正文"),
            *tables,
        ]

    def pdf(self, source: Path, directory: Path) -> list[ReferencePart]:
        import fitz

        parts = []
        with fitz.open(source) as document:
            if document.needs_pass:
                raise ValueError("PDF 已加密，不能读取。")
            for number, page in enumerate(document, 1):
                locator = f"第 {number} 页"
                text = page.get_text(sort=True)
                if text.strip():
                    parts.append(ReferencePart.from_text(text, locator=locator))
                for table in page.find_tables().tables:
                    parts.append(
                        ReferencePart.from_text(
                            table.extract(), locator=locator + " 表格"
                        )
                    )
                if page.get_images() or not text.strip() or page.get_drawings():
                    target = directory / f"page-{number}.png"
                    page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4)).save(target)
                    parts.append(
                        ReferencePart(
                            kind="image",
                            locator=locator,
                            path=target,
                            media_type="image/png",
                        )
                    )
        return parts

    def word(self, source: Path, directory: Path) -> list[ReferencePart]:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn

        document = Document(source)
        parts, included = [], set()

        def append_image(relation_id, locator):
            relation = document.part.rels[relation_id]
            if relation.is_external or "image" not in relation.reltype:
                return
            image = relation.target_part
            target = (
                directory / f"word-image-{len(parts)}{Path(str(image.partname)).suffix}"
            )
            target.write_bytes(image.blob)
            parts.append(
                ReferencePart(
                    kind="image",
                    locator=locator,
                    path=target,
                    media_type=image.content_type,
                )
            )
            included.add(relation_id)

        for index, block in enumerate(document.iter_inner_content(), 1):
            if isinstance(block, Paragraph):
                text = block.text
            elif isinstance(block, Table):
                text = json.dumps(
                    [[cell.text for cell in row.cells] for row in block.rows],
                    ensure_ascii=False,
                )
            else:
                continue
            if text.strip():
                parts.append(ReferencePart.from_text(text, locator=f"内容块 {index}"))
            for blip in block._element.iter(qn("a:blip")):
                if relation_id := blip.get(qn("r:embed")):
                    append_image(relation_id, f"内容块 {index} / 图片")
        for relation_id in document.part.rels:
            if relation_id not in included:
                append_image(relation_id, "文档其他图片")
        return parts

    def excel(self, source: Path, directory: Path) -> list[ReferencePart]:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        formulas = load_workbook(source, read_only=True, data_only=False)
        cached = load_workbook(source, read_only=True, data_only=True)
        parts = []
        try:
            for sheet in formulas:
                rows = []
                for row_number, (source_row, cached_row) in enumerate(
                    zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1
                ):
                    row = []
                    for column, (cell, value) in enumerate(
                        zip(source_row, cached_row), 1
                    ):
                        item = dict(
                            cell=f"{get_column_letter(column)}{row_number}",
                            value=value.value,
                        )
                        if cell.data_type == "f":
                            item["formula"] = cell.value
                        row.append(item)
                    rows.append(row)
                parts.append(
                    ReferencePart.from_text(rows, locator=f"工作表 {sheet.title}")
                )
        finally:
            formulas.close()
            cached.close()
        return parts

    def slides(self, source: Path, directory: Path) -> list[ReferencePart]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        parts = []
        for number, slide in enumerate(Presentation(source).slides, 1):

            def read_shapes(shapes):
                for shape in shapes:
                    locator = f"幻灯片 {number} / 形状 {shape.shape_id}"
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        read_shapes(shape.shapes)
                    if shape.has_text_frame and shape.text.strip():
                        parts.append(
                            ReferencePart.from_text(shape.text, locator=locator)
                        )
                    if shape.has_table:
                        parts.append(
                            ReferencePart.from_text(
                                [
                                    [cell.text for cell in row.cells]
                                    for row in shape.table.rows
                                ],
                                locator=locator,
                            )
                        )
                    if shape.has_chart:
                        values = [
                            dict(name=series.name, values=list(series.values))
                            for series in shape.chart.series
                        ]
                        parts.append(
                            ReferencePart.from_text(values, locator=locator + " 图表")
                        )
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        target = (
                            directory
                            / f"slide-{number}-image-{shape.shape_id}.{shape.image.ext}"
                        )
                        target.write_bytes(shape.image.blob)
                        parts.append(
                            ReferencePart(
                                kind="image",
                                locator=locator,
                                path=target,
                                media_type=shape.image.content_type,
                            )
                        )

            read_shapes(slide.shapes)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    parts.append(
                        ReferencePart.from_text(
                            notes.text, locator=f"幻灯片 {number} 备注"
                        )
                    )
        return parts
