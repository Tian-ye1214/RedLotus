"""Immutable reference resources, document parsing, and browser access."""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import mimetypes
import os
import shutil
import tempfile
from pathlib import Path
from typing import Literal
from urllib.parse import urlsplit

from filelock import AsyncFileLock
from pydantic import BaseModel, ConfigDict, Field
from pydantic_ai import BinaryContent

from redlotus.runtime.config import get_env
from redlotus.runtime.network import ModelInputPolicy
from redlotus.runtime.resources import (
    WorkspaceContext,
    atomic_write_bytes,
    atomic_write_json,
    finish_file_io,
    references_dir,
    runtime_dir,
    user_data_dir,
)
from redlotus.tools.execution import run_subprocess


class ReferencePart(BaseModel):
    model_config = ConfigDict(frozen=True)
    kind: Literal["text", "image", "video", "audio"]
    locator: str = ""
    text: str = ""
    path: Path | None = None
    media_type: str = ""

    @classmethod
    def from_text(cls, value, *, locator=""):
        text = (
            value
            if isinstance(value, str)
            else json.dumps(value, ensure_ascii=False, default=str)
        )
        return cls(kind="text", text=text, locator=locator)


class ReferenceFile(BaseModel):
    model_config = ConfigDict(frozen=True)
    id: str
    project_id: str
    name: str
    source: str
    media_type: str
    byte_size: int
    sha256: str
    snapshot: Path
    parts: list[ReferencePart] = Field(default_factory=list)
    parser_version: int = 1

    def to_prompt(self) -> list:
        """Present reference identity, coverage and complete content without extra instructions."""
        header = self.model_dump(mode="json", exclude={"parts"})
        header["content_provided"] = bool(self.parts)
        header["coverage"] = [
            {"part": index, "locator": part.locator, "kind": part.kind}
            for index, part in enumerate(self.parts)
        ]
        content = [json.dumps({"reference_file": header}, ensure_ascii=False)]
        for index, part in enumerate(self.parts):
            label = json.dumps({
                "reference_id": self.id, "name": self.name,
                "part": index, "locator": part.locator, "kind": part.kind,
            }, ensure_ascii=False)
            if part.kind == "text":
                content.append(label + "\n" + part.text)
            else:
                content.append(label)
                content.append(BinaryContent(
                    data=part.path.read_bytes(),
                    media_type=part.media_type,
                    identifier=f"{self.id}-{index}",
                ))
        content.append(json.dumps({"reference_end": self.id}))
        return content

    def manifest(self) -> dict:
        return self.model_dump(mode="json")


class OfficeConverter:
    """A private headless Office profile, with cancellable process-tree ownership."""

    @staticmethod
    def cleanup_profile(profile: Path) -> None:
        """Remove an Office profile without the Windows MAX_PATH fallback limit."""
        path = str(profile.resolve())
        if os.name == "nt":
            path = (
                "\\\\?\\UNC\\" + path[2:]
                if path.startswith("\\\\")
                else "\\\\?\\" + path
            )
        shutil.rmtree(path)

    @staticmethod
    def executable() -> str:
        candidates = [
            get_env("LIBREOFFICE_PATH", warn=False),
            shutil.which("soffice.com"),
            shutil.which("soffice"),
            str(
                Path(os.environ.get("ProgramFiles", "C:/Program Files"))
                / "LibreOffice/program/soffice.com"
            ),
            str(user_data_dir() / "tools/libreoffice/program/soffice.com"),
        ]
        for value in candidates:
            if value and Path(value).is_file():
                return value
        raise ValueError(
            "DOC/PPT 转换需要 LibreOffice；请安装或设置 LIBREOFFICE_PATH。"
        )

    async def convert(self, source: Path, target_format: str, directory: Path) -> Path:
        executable = self.executable()
        directory.mkdir(parents=True, exist_ok=True)
        # LibreOffice still uses Windows APIs with limited path lengths; keep its private
        # profile and working copies short, then publish the result to the reference store.
        scratch = runtime_dir() / "office"
        scratch.mkdir(parents=True, exist_ok=True)
        profile_root = Path(tempfile.mkdtemp(prefix="rl-office-", dir=scratch))
        try:
            user = profile_root / "user"
            user.mkdir()
            working = profile_root / "document" / source.name
            working.parent.mkdir()
            shutil.copyfile(source, working)
            (user / "registrymodifications.xcu").write_text(
                '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
                '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>',
                encoding="utf-8",
            )
            args = [
                executable,
                "-env:UserInstallation=" + profile_root.as_uri(),
                "--headless",
                "--nologo",
                "--nodefault",
                "--norestore",
                "--convert-to",
                target_format,
                "--outdir",
                str(profile_root),
                str(working),
            ]
            result = await run_subprocess(
                args, shell=False, cwd=str(profile_root), timeout=120
            )
            output = profile_root / (source.stem + "." + target_format.split(":")[0])
            if result.returncode != 0 or not output.is_file():
                raise ValueError(
                    f"Office 转换失败 ({source.name})：{result.to_text()}"
                )
            target = directory / output.name
            shutil.copyfile(output, target)
            return target
        finally:
            self.cleanup_profile(profile_root)


class DocumentReader:
    """Read format structure once; never execute referenced content."""

    async def read(self, source: Path, directory: Path) -> list[ReferencePart]:
        extension = source.suffix.lower()
        if extension in (".doc", ".ppt", ".xls"):
            modern = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}[extension]
            source = await OfficeConverter().convert(
                source, modern, directory / "converted"
            )
            extension = "." + modern
        readers = {
            ".pdf": self.pdf,
            ".docx": self.word,
            ".pptx": self.slides,
            ".xlsx": self.excel,
            ".csv": self.csv,
            ".html": self.html,
            ".htm": self.html,
        }
        return await finish_file_io(
            asyncio.to_thread(readers.get(extension, self.text), source, directory)
        )

    @staticmethod
    def decode(data: bytes) -> str:
        for encoding in (
            "utf-8-sig",
            "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "gb18030",
        ):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                pass
        raise ValueError("无法识别引用文件的文本编码。")

    def text(self, source: Path, directory: Path) -> list[ReferencePart]:
        text = self.decode(source.read_bytes())
        if "\x00" in text:
            raise ValueError(f"没有配置此二进制格式的解析器：{source.suffix}")
        return [ReferencePart.from_text(text)]

    def csv(self, source: Path, directory: Path) -> list[ReferencePart]:
        text = self.decode(source.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        return [ReferencePart.from_text(rows, locator="CSV rows and columns")]

    def html(self, source: Path, directory: Path) -> list[ReferencePart]:
        from lxml import etree, html

        root = html.fromstring(
            self.decode(source.read_bytes()).encode("utf-8"),
            parser=html.HTMLParser(encoding="utf-8"),
        )
        etree.strip_elements(root, "script", "style", with_tail=False)
        tables = []
        for number, table in enumerate(root.xpath("self::table | .//table"), 1):
            rows = [
                [cell.text_content().strip() for cell in row.xpath("./th | ./td")]
                for row in table.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
            ]
            tables.append(ReferencePart.from_text(rows, locator=f"HTML table {number}"))
            table.drop_tree()
        for link in root.xpath(".//a[@href]"):
            link.tail = f" ({link.get('href')})" + (link.tail or "")
        for element in root.iter():
            if element.tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
                element.text = "#" * int(element.tag[1]) + " " + (element.text or "")
            if element.tag in (
                "p",
                "div",
                "section",
                "li",
                "br",
                "h1",
                "h2",
                "h3",
                "h4",
                "h5",
                "h6",
            ):
                element.tail = "\n" + (element.tail or "")
        return [
            ReferencePart.from_text(root.text_content().strip(), locator="HTML body"),
            *tables,
        ]

    def pdf(self, source: Path, directory: Path) -> list[ReferencePart]:
        import fitz

        parts = []
        with fitz.open(source) as document:
            if document.needs_pass:
                raise ValueError("PDF 已加密，不能读取。")
            for number, page in enumerate(document, 1):
                locator = f"Page {number}"
                text = page.get_text(sort=True)
                if text.strip():
                    parts.append(ReferencePart.from_text(text, locator=locator))
                for table in page.find_tables().tables:
                    parts.append(
                        ReferencePart.from_text(
                            table.extract(), locator=f"{locator}, table"
                        )
                    )
                if page.get_images() or not text.strip() or page.get_drawings():
                    target = directory / f"page-{number}.png"
                    page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4)).save(target)
                    parts.append(
                        ReferencePart(
                            kind="image",
                            locator=locator,
                            path=target,
                            media_type="image/png",
                        )
                    )
        return parts

    def word(self, source: Path, directory: Path) -> list[ReferencePart]:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(source)
        parts, included = [], set()

        def append_image(relation_id, locator):
            relation = document.part.rels[relation_id]
            if relation.is_external or "image" not in relation.reltype:
                return
            image = relation.target_part
            target = (
                directory / f"word-image-{len(parts)}{Path(str(image.partname)).suffix}"
            )
            target.write_bytes(image.blob)
            parts.append(
                ReferencePart(
                    kind="image",
                    locator=locator,
                    path=target,
                    media_type=image.content_type,
                )
            )
            included.add(relation_id)

        for index, block in enumerate(document.iter_inner_content(), 1):
            if isinstance(block, Paragraph):
                text = block.text
            elif isinstance(block, Table):
                text = json.dumps(
                    [[cell.text for cell in row.cells] for row in block.rows],
                    ensure_ascii=False,
                )
            else:
                continue
            if text.strip():
                parts.append(ReferencePart.from_text(text, locator=f"Document block {index}"))
            for blip in block._element.iter(qn("a:blip")):
                if relation_id := blip.get(qn("r:embed")):
                    append_image(relation_id, f"Document block {index}, image")
        for relation_id in document.part.rels:
            if relation_id not in included:
                append_image(relation_id, "Document image")
        return parts

    def excel(self, source: Path, directory: Path) -> list[ReferencePart]:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        formulas = load_workbook(source, read_only=True, data_only=False)
        cached = load_workbook(source, read_only=True, data_only=True)
        parts = []
        try:
            for sheet in formulas:
                rows = []
                for row_number, (source_row, cached_row) in enumerate(
                    zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1
                ):
                    row = []
                    for column, (cell, value) in enumerate(
                        zip(source_row, cached_row), 1
                    ):
                        item = dict(
                            cell=f"{get_column_letter(column)}{row_number}",
                            value=value.value,
                        )
                        if cell.data_type == "f":
                            item["formula"] = cell.value
                        row.append(item)
                    rows.append(row)
                parts.append(
                    ReferencePart.from_text(rows, locator=f"Sheet {sheet.title}")
                )
        finally:
            formulas.close()
            cached.close()
        return parts

    def slides(self, source: Path, directory: Path) -> list[ReferencePart]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        parts = []
        for number, slide in enumerate(Presentation(source).slides, 1):

            def read_shapes(shapes):
                for shape in shapes:
                    locator = f"Slide {number}, shape {shape.shape_id}"
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        read_shapes(shape.shapes)
                    if shape.has_text_frame and shape.text.strip():
                        parts.append(
                            ReferencePart.from_text(shape.text, locator=locator)
                        )
                    if shape.has_table:
                        parts.append(
                            ReferencePart.from_text(
                                [
                                    [cell.text for cell in row.cells]
                                    for row in shape.table.rows
                                ],
                                locator=locator,
                            )
                        )
                    if shape.has_chart:
                        values = [
                            dict(name=series.name, values=list(series.values))
                            for series in shape.chart.series
                        ]
                        parts.append(
                            ReferencePart.from_text(values, locator=f"{locator}, chart")
                        )
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        target = (
                            directory
                            / f"slide-{number}-image-{shape.shape_id}.{shape.image.ext}"
                        )
                        target.write_bytes(shape.image.blob)
                        parts.append(
                            ReferencePart(
                                kind="image",
                                locator=locator,
                                path=target,
                                media_type=shape.image.content_type,
                            )
                        )

            read_shapes(slide.shapes)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    parts.append(
                        ReferencePart.from_text(
                            notes.text, locator=f"Slide {number}, notes"
                        )
                    )
        return parts


def reference_message_data(value, *, restore=False, workspace=None):
    """Keep native attachment bytes in immutable snapshots, with links in session JSON."""
    import base64

    from redlotus.runtime import resources as paths
    from redlotus.runtime.resources import file_lock

    if isinstance(value, list):
        return [reference_message_data(item, restore=restore, workspace=workspace) for item in value]
    if not isinstance(value, dict):
        return value
    if value.get("kind") != "binary" or "data" not in value:
        return {key: reference_message_data(item, restore=restore, workspace=workspace) for key, item in value.items()}
    root = paths.references_dir(workspace).resolve()
    if restore:
        if not isinstance(value["data"], dict):
            return value
        descriptor = value["data"]
        snapshot = Path(descriptor["snapshot"]).resolve()
        if not snapshot.is_relative_to(root):
            raise ValueError("引用快照不属于配置的引用目录")
        data = snapshot.read_bytes()
        if hashlib.sha256(data).hexdigest() != descriptor["sha256"]:
            raise ValueError(f"引用快照校验失败：{snapshot}")
        return {**value, "data": base64.urlsafe_b64encode(data).decode()}
    data = base64.urlsafe_b64decode(value["data"])
    digest = hashlib.sha256(data).hexdigest()
    identifier = str(value.get("identifier", ""))
    reference_id, _, part = identifier.partition("-")
    manifest = root / "manifests" / f"{reference_id}.json"
    snapshot = None
    if part.isdigit() and manifest.is_file():
        parts = json.loads(manifest.read_text(encoding="utf-8"))["parts"]
        snapshot = Path(parts[int(part)]["path"])
    if snapshot is None:
        directory = root / "blobs" / digest[:32]
        snapshot = next(directory.glob("source.*"), directory / "source.bin")
        with file_lock(snapshot):
            if not snapshot.exists():
                atomic_write_bytes(snapshot, data)
    return {**value, "data": {"snapshot": str(snapshot), "sha256": digest}}


class ReferenceStore:
    PARSER_VERSION = 3

    def __init__(self, workspace: WorkspaceContext, root: Path | None = None):
        self.workspace = workspace
        self.root = root or references_dir(workspace)

    async def prepare_message(self, message):
        import base64
        import mimetypes

        from pydantic_ai import BinaryContent, ImageUrl, VideoUrl

        from redlotus.runtime.network import ModelInputPolicy

        policy = ModelInputPolicy.for_role("coordinator")
        references = list(message.references)
        remote_images = {}
        for index, item in enumerate(message.attachments):
            if isinstance(item, BinaryContent):
                ref = await self.import_binary(item, source=f"attachment:{index}", policy=policy)
            elif isinstance(item, ImageUrl) and urlsplit(item.url).scheme in ("http", "https"):
                remote_images.setdefault(item.url, item)
                continue
            elif isinstance(item, (ImageUrl, VideoUrl)):
                if item.url.startswith("data:"):
                    header, encoded = item.url.split(",", 1)
                    mime = header[5:].split(";")[0]
                    name = "attachment" + (mimetypes.guess_extension(mime) or ".bin")
                    ref = await self.import_bytes(
                        base64.b64decode(encoded),
                        name=name,
                        source=f"attachment:{index}",
                        policy=policy,
                    )
                else:
                    ref = await self.import_url(item.url, policy=policy)
            else:
                raise ValueError(f"Unsupported attachment type: {type(item).__name__}")
            references.append(ref)
        references = list({ref.id: ref for ref in references}.values())
        count = len(references) + len(remote_images)
        if count > policy.max_files:
            raise ValueError(f"最多引用 {policy.max_files} 个文件，本次为 {count} 个。")
        # Remote image bytes are fetched and checked by the receiving gateway.
        policy.check([ref.byte_size for ref in references])
        message.references, message.attachments = references, list(remote_images.values())

    async def import_binary(self, item, *, source, policy):
        """Reuse registered native media, or capture an unregistered attachment once."""
        try:
            registered = self.load(item.identifier[:32])
        except (ValueError, FileNotFoundError):
            registered = None
        if registered is not None:
            return await self.parse(registered)
        name = item.identifier or "media"
        if not Path(name).suffix:
            name += mimetypes.guess_extension(item.media_type) or ".bin"
        return await self.import_bytes(item.data, name=name, source=source, policy=policy)

    async def capture_file(
        self, path: Path, *, policy: ModelInputPolicy
    ) -> ReferenceFile:
        policy.check([path.stat().st_size])
        data = await asyncio.to_thread(path.read_bytes)
        return await self.capture_bytes(
            data, name=path.name, source=str(path), policy=policy
        )

    async def import_file(
        self, path: Path, *, policy: ModelInputPolicy
    ) -> ReferenceFile:
        return await self.parse(await self.capture_file(path, policy=policy))

    async def import_url(
        self, url: str, *, policy: ModelInputPolicy, media_type: str = ""
    ) -> ReferenceFile:
        from urllib.parse import unquote, urlsplit

        import httpx

        from redlotus.runtime.network import get_client

        if urlsplit(url).scheme not in ("https", "http"):
            raise ValueError("Remote references require HTTP(S)")
        client = get_client(
            "reference_download",
            lambda: httpx.AsyncClient(timeout=60, follow_redirects=True),
        )
        async with client.stream("GET", url) as response:
            response.raise_for_status()
            data = bytearray()
            async for chunk in response.aiter_bytes():
                data.extend(chunk)
                policy.check([len(data)])
            mime = media_type or response.headers.get("content-type", "").split(";")[0]
        name = Path(unquote(urlsplit(url).path)).name or "attachment"
        if not Path(name).suffix:
            name += mimetypes.guess_extension(mime) or ".bin"
        return await self.import_bytes(
            bytes(data), name=name, source=url, policy=policy
        )

    async def import_bytes(
        self, data: bytes, *, name: str, source: str, policy: ModelInputPolicy
    ) -> ReferenceFile:
        reference = await self.capture_bytes(
            data, name=name, source=source, policy=policy
        )
        return await self.parse(reference)

    async def capture_bytes(
        self, data: bytes, *, name: str, source: str, policy: ModelInputPolicy
    ) -> ReferenceFile:
        policy.check([len(data)])
        media_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
        if media_type.startswith("image/"):
            from PIL import Image

            with Image.open(io.BytesIO(data)) as picture:
                media_type = Image.MIME[picture.format]
        digest = hashlib.sha256(data).hexdigest()
        identity = hashlib.sha256(
            f"{self.workspace.project_id}\0{os.path.normcase(source)}\0{digest}".encode()
        ).hexdigest()[:32]
        directory = self.root / "blobs" / digest[:32]
        directory.mkdir(parents=True, exist_ok=True)
        snapshot = directory / ("source" + Path(name).suffix.lower())
        async with AsyncFileLock(directory / ".build.lock", run_in_executor=False):
            if not snapshot.exists():
                await finish_file_io(
                    asyncio.to_thread(atomic_write_bytes, snapshot, data)
                )
        return ReferenceFile(
            id=identity,
            project_id=self.workspace.project_id,
            name=name,
            source=source,
            media_type=media_type,
            byte_size=len(data),
            sha256=digest,
            snapshot=snapshot,
        )

    async def parse(self, reference: ReferenceFile) -> ReferenceFile:
        manifest = self.root / "manifests" / f"{reference.id}.json"
        manifest.parent.mkdir(parents=True, exist_ok=True)
        async with AsyncFileLock(str(manifest) + ".lock", run_in_executor=False):
            if manifest.is_file():
                cached = ReferenceFile.model_validate_json(
                    manifest.read_text(encoding="utf-8")
                )
                if cached.parser_version == self.PARSER_VERSION:
                    return cached
            snapshot = reference.snapshot
            directory = snapshot.parent
            parts_path = directory / (
                f"parts-v{self.PARSER_VERSION}" + snapshot.suffix + ".json"
            )
            async with AsyncFileLock(directory / ".build.lock", run_in_executor=False):
                if parts_path.is_file():
                    parts = [
                        ReferencePart.model_validate(item)
                        for item in json.loads(parts_path.read_text(encoding="utf-8"))
                    ]
                else:
                    if reference.media_type.startswith(("image/", "video/", "audio/")):
                        parts = await finish_file_io(
                            asyncio.to_thread(self._media_parts, reference)
                        )
                    else:
                        parts = await DocumentReader().read(snapshot, directory)
                    await finish_file_io(
                        asyncio.to_thread(
                            atomic_write_json,
                            parts_path,
                            [part.model_dump(mode="json") for part in parts],
                        )
                    )
            prepared = reference.model_copy(
                update={"parts": parts, "parser_version": self.PARSER_VERSION}
            )
            await finish_file_io(
                asyncio.to_thread(atomic_write_json, manifest, prepared.manifest())
            )
            return prepared

    @staticmethod
    def _media_parts(reference: ReferenceFile) -> list[ReferencePart]:
        kind = reference.media_type.split("/")[0]
        path, media_type = reference.snapshot, reference.media_type
        if kind == "image":
            from PIL import Image

            with Image.open(path) as picture:
                media_type = Image.MIME[picture.format]
                picture.verify()
            if media_type not in ("image/png", "image/jpeg", "image/webp", "image/gif"):
                with Image.open(path) as picture:
                    path = path.parent / "image.png"
                    picture.save(path)
                media_type = "image/png"
        elif kind == "video":
            with path.open("rb") as stream:
                header = stream.read(16)
            if not (
                header[4:8] == b"ftyp" or header.startswith((b"RIFF", b"\x1aE\xdf\xa3"))
            ):
                raise ValueError("视频容器无效或未识别，不能作为原生视频提交。")
        return [
            ReferencePart(kind=kind, path=path, media_type=media_type, locator="Original file")
        ]

    def load(self, reference_id: str) -> ReferenceFile:
        if len(reference_id) != 32 or any(
            c not in "0123456789abcdef" for c in reference_id
        ):
            raise ValueError("引用文件 ID 无效。")
        path = self.root / "manifests" / f"{reference_id}.json"
        return ReferenceFile.model_validate_json(path.read_text(encoding="utf-8"))

    async def read_reference(self, reference_id: str):
        """Read the complete immutable snapshot of a registered project reference.

        Use the supplied reference content directly when it is already in context. Use
        this tool when the registered snapshot must be retrieved again; read_file reads
        the current disk version, and extract_text registers a project document.

        Args:
            reference_id: The reference ID supplied with the attachment or memory record.

        Returns:
            The original reference identity, parsed text and native media, or an error."""
        from pydantic_ai import ToolReturn

        try:
            reference = await asyncio.to_thread(self.load, reference_id)
            if reference.project_id != self.workspace.project_id:
                return "Error: Reference belongs to another project; retrieve its authorized memory record instead."
            reference = await self.parse(reference)
        except (OSError, ValueError) as exc:
            return f"Error reading reference '{reference_id}': {exc}"
        return ToolReturn(
            return_value=f"Read reference {reference.name} ({reference.id})",
            content=await asyncio.to_thread(reference.to_prompt),
        )
